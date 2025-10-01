"""
Módulo para el procesamiento del 1_Reporte_Llamadas - Versión Optimizada
Incluye procesamiento para Reporte Isabel y Reporte VOIP
"""
import pandas as pd
import io
from datetime import datetime
from typing import List, Tuple, Optional, Dict
from flask import request, send_file
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment

# ==============================================================================
# CONSTANTES Y CONFIGURACIÓN
# ==============================================================================

# Extensiones permitidas para archivos de llamadas
ALLOWED_EXTENSIONS_CALLS = {'csv', 'xlsx', 'xls'}

# Diccionario para nombres de meses en español
MESES_ESPANOL = {
    1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 
    5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
    9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
}

# Estado de llamada efectiva
ESTADO_EFECTIVA = 'ANSWERED'

# Estados válidos que se deben contar en el total VOIP
ESTADOS_VALIDOS_VOIP = {'answered', 'busy', 'no_answer', 'normal', 'out_area', 'offline'}

# Columnas requeridas para Reporte VOIP
COLUMNAS_REQUERIDAS_VOIP = ['First Call Agent', 'Ring Type', 'Begin Time']

# Mapeo de columnas chino-inglés
COLUMNAS_CHINAS_MAPEO = {
    '外呼人員': 'First Call Agent',
    '状态': 'Ring Type',
    '开始时间': 'Begin Time'
}

# Configuración Excel
TABLA_ESTILO = "TableStyleMedium12"
FORMATO_NUMERO = '0'
FORMATO_HORA = 'h:mm:ss AM/PM'

# Configuración de columnas
COLUMNAS_CONFIG_ISABEL = {
    "Número Extensión": {"width": 18, "format": FORMATO_NUMERO},
    "Total Llamadas": {"width": 16, "format": FORMATO_NUMERO},
    "Llamadas Efectivas": {"width": 18, "format": FORMATO_NUMERO},
    "Última Llamada": {"width": 16, "format": FORMATO_HORA}
}

COLUMNAS_CONFIG_VOIP = {
    "Extensión": {"width": 20, "format": FORMATO_NUMERO},
    "Total": {"width": 12, "format": FORMATO_NUMERO},
    "Ultima Llamada": {"width": 16, "format": FORMATO_HORA}
}

# ==============================================================================
# FUNCIONES
# ==============================================================================

def allowed_file_calls(filename: str) -> bool:
    """
    Verifica si el archivo tiene una extensión válida para procesamiento de llamadas
    
    Args:
        filename: Nombre del archivo a validar
        
    Returns:
        bool: True si la extensión es válida (CSV, XLSX, XLS), False caso contrario
    """
    if not filename or '.' not in filename:
        return False
    
    extension = filename.rsplit('.', 1)[1].lower()
    return extension in ALLOWED_EXTENSIONS_CALLS

def validar_archivos_entrada(archivos_key: str = 'files') -> Tuple[Optional[List], Optional[str], Optional[int]]:
    """
    Valida que se hayan subido archivos válidos (CSV o Excel)
    
    Args:
        archivos_key: Clave para obtener archivos del request (default: 'files')
        
    Returns:
        Tuple con (archivos_validos, mensaje_error, codigo_http)
        - Si éxito: (List[FileStorage], None, None)
        - Si error: (None, str, int)
        
    Raises:
        None - Todos los errores se manejan como retorno de tuple
    """
    try:
        files = request.files.getlist(archivos_key)
        
        # Verificar que existan archivos
        if not files:
            return None, "Error: Debes subir al menos un archivo (CSV o Excel).", 400
        
        # Validar cada archivo individualmente
        archivos_validos = []
        for file in files:
            if not file:
                continue  # Saltar archivos vacíos
                
            filename = getattr(file, 'filename', None)
            if not filename:
                return None, "Error: Se encontró un archivo sin nombre.", 400
                
            if not allowed_file_calls(filename):
                return None, f"Error: El archivo '{filename}' no es válido. Extensiones permitidas: CSV, XLSX, XLS.", 400
                
            archivos_validos.append(file)
        
        # Verificar que quedaron archivos válidos después del filtrado
        if not archivos_validos:
            return None, "Error: No se encontraron archivos válidos para procesar.", 400
        
        return archivos_validos, None, None
        
    except Exception as e:
        return None, f"Error interno al validar archivos: {str(e)}", 500

def leer_archivo_datos(file) -> pd.DataFrame:
    """
    Lee un archivo (CSV o Excel) y retorna un DataFrame con validaciones robustas
    
    Args:
        file: Objeto FileStorage con el archivo a leer
        
    Returns:
        pd.DataFrame: DataFrame con los datos del archivo
        
    Raises:
        ValueError: Si el archivo no es válido o no se puede leer
        FileNotFoundError: Si el archivo no existe
        pd.errors.EmptyDataError: Si el archivo está vacío
    """
    if not file:
        raise ValueError("El objeto archivo no puede ser None")
    
    filename = getattr(file, 'filename', None)
    if not filename:
        raise ValueError("El archivo no tiene nombre válido")
    
    filename_lower = filename.lower()
    
    try:
        # Configuración común para lectura
        read_kwargs = {
            'dtype': str,  # Leer todo como string inicialmente para evitar conversiones automáticas
        }
        
        # Leer según el tipo de archivo
        if filename_lower.endswith('.csv'):
            # Configuración específica para CSV
            csv_kwargs = {
                **read_kwargs,
                'encoding': 'utf-8',  # Encoding por defecto
                'sep': ',',           # Separador por defecto
            }
            
            try:
                df = pd.read_csv(file, **csv_kwargs)
            except UnicodeDecodeError:
                # Intentar con encoding alternativo
                file.seek(0)  # Resetear posición del archivo
                csv_kwargs['encoding'] = 'latin-1'
                df = pd.read_csv(file, **csv_kwargs)
                
        elif filename_lower.endswith(('.xlsx', '.xls')):
            # Leer archivo Excel
            df = pd.read_excel(file, **read_kwargs)
            
        else:
            raise ValueError(f"Tipo de archivo no soportado: {filename}. "
                           f"Extensiones permitidas: .csv, .xlsx, .xls")
        
        # Validaciones del DataFrame resultante
        if df.empty:
            raise pd.errors.EmptyDataError(f"El archivo '{filename}' está vacío o no contiene datos válidos")
        
        # Limpiar nombres de columnas (eliminar espacios extra)
        df.columns = df.columns.str.strip()
        
        return df
        
    except pd.errors.EmptyDataError:
        raise
    except Exception as e:
        raise ValueError(f"Error al leer el archivo '{filename}': {str(e)}")

def aplicar_mapeo_columnas_chinas(df: pd.DataFrame) -> pd.DataFrame:
    """
    Aplica el mapeo de columnas chinas a inglés
    
    Args:
        df: DataFrame con columnas potencialmente en chino
        
    Returns:
        DataFrame con columnas mapeadas
    """
    df_mapped = df.copy()
    
    for col_china, col_ingles in COLUMNAS_CHINAS_MAPEO.items():
        if col_china in df_mapped.columns:
            df_mapped = df_mapped.rename(columns={col_china: col_ingles})
    
    return df_mapped

def generar_nombre_hoja_fecha(day) -> str:
    """
    Genera el nombre de hoja con formato DD-MM-YYYY (válido para Excel)
    
    Args:
        day: Fecha para generar el nombre
        
    Returns:
        Nombre de hoja formateado como DD-MM-YYYY (Excel no permite / en nombres de hoja)
    """
    dt = pd.to_datetime(day)
    return f"{dt.day:02d}-{dt.month:02d}-{dt.year}"

def configurar_hoja_excel(worksheet, df: pd.DataFrame, columnas_config: Dict, tabla_prefix: str) -> None:
    """
    Configura una hoja de Excel con formato de tabla y estilos
    
    Args:
        worksheet: Hoja de trabajo de openpyxl
        df: DataFrame con los datos
        columnas_config: Configuración de columnas
        tabla_prefix: Prefijo para el nombre de la tabla
    """
    num_rows, num_cols = df.shape
    
    if num_rows == 0 or num_cols == 0:
        return
    
    # Crear tabla Excel
    import re
    from openpyxl.utils import get_column_letter
    
    clean_title = re.sub(r'[^a-zA-Z0-9_]', '_', worksheet.title)
    table_name = f"{tabla_prefix}{clean_title}"[:31]
    end_col = get_column_letter(num_cols)
    table_range = f"A1:{end_col}{num_rows + 1}"
    
    table = Table(displayName=table_name, ref=table_range)
    table.tableStyleInfo = TableStyleInfo(
        name=TABLA_ESTILO, showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    worksheet.add_table(table)
    
    # Aplicar formato a columnas
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    for col_obj in worksheet.columns:
        column_letter = col_obj[0].column_letter
        column_name = worksheet[f"{column_letter}1"].value
        
        # Configurar ancho de columna
        if column_name in columnas_config:
            worksheet.column_dimensions[column_letter].width = columnas_config[column_name]["width"]
        else:
            max_length = max(len(str(cell.value)) for cell in col_obj if cell.value)
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
        
        # Aplicar alineación y formato
        for row_num, cell in enumerate(col_obj, 1):
            cell.alignment = center_alignment
            
            if row_num > 1 and column_name in columnas_config:  # No aplicar formato a encabezados
                cell.number_format = columnas_config[column_name]["format"]

# ==============================================================================
# FUNCIONES DE PROCESAMIENTO DE DATOS
# ==============================================================================

def procesar_datos_csv(*files) -> Tuple[pd.DataFrame, str]:
    """
    Procesa y combina múltiples archivos CSV/Excel para generar el Reporte Isabel
    
    Args:
        *files: Objetos FileStorage con los archivos a procesar
        
    Returns:
        Tuple[pd.DataFrame, str]: DataFrame procesado y nombre de la columna de fecha encontrada
        
    Raises:
        ValueError: Si no se encuentra columna de fecha o hay problemas con los datos
    """
    dataframes = []
    
    # Leer y combinar todos los archivos
    for file in files:
        df = leer_archivo_datos(file)
        dataframes.append(df)
    
    # Combinar todos los DataFrames
    df = pd.concat(dataframes, ignore_index=True)
    
    # Encontrar columna de fecha automáticamente
    date_col_name = next(
        (col for col in df.columns if 'fecha' in col.lower()), 
        None
    )
    if not date_col_name:
        available_cols = ', '.join(df.columns)
        raise ValueError(
            f"No se encontró columna de fecha en los archivos. "
            f"Columnas disponibles: {available_cols}"
        )
    
    # Limpiar datos de fuente (extensión)
    df['Fuente'] = pd.to_numeric(df['Fuente'], errors='coerce')
    df.dropna(subset=['Fuente'], inplace=True)
    df['Fuente'] = df['Fuente'].astype(int)
    df = df[df['Fuente'] <= 30000]
    
    # Limpiar datos de fecha
    df[date_col_name] = pd.to_datetime(df[date_col_name], errors='coerce')
    df.dropna(subset=[date_col_name], inplace=True)
    df['day_key'] = df[date_col_name].dt.date

    return df, date_col_name

def procesar_datos_agentes(*files) -> pd.DataFrame:
    """
    Procesa archivos CSV/Excel con datos de agentes para generar el Reporte VOIP
    Incluye traducción automática de columnas del chino al inglés
    
    Args:
        *files: Objetos FileStorage con los archivos a procesar
        
    Returns:
        pd.DataFrame: DataFrame procesado con datos de agentes limpios
        
    Raises:
        ValueError: Si faltan columnas requeridas o hay problemas con los datos
    """
    dataframes = []
    
    # Leer y combinar archivos
    for file in files:
        df = leer_archivo_datos(file)
        df = aplicar_mapeo_columnas_chinas(df)
        dataframes.append(df)
    
    df = pd.concat(dataframes, ignore_index=True)
    
    # Verificar columnas requeridas para VOIP
    missing_columns = [
        col for col in COLUMNAS_REQUERIDAS_VOIP 
        if col not in df.columns
    ]
    if missing_columns:
        available_columns = ', '.join(df.columns)
        missing_str = ', '.join(missing_columns)
        raise ValueError(
            f"Faltan columnas requeridas: {missing_str}. "
            f"Columnas disponibles: {available_columns}"
        )
    
    # Limpiar datos
    df.dropna(subset=['First Call Agent'], inplace=True)
    df = df[df['First Call Agent'].astype(str).str.strip() != '']
    
    # Formatear nombres de agentes usando str.title() directamente
    df['First Call Agent'] = df['First Call Agent'].astype(str).str.title()
    
    df['Ring Type'] = df['Ring Type'].astype(str).str.lower().str.strip()
    
    # Filtrar solo estados válidos y procesar fechas
    df = df[df['Ring Type'].isin(ESTADOS_VALIDOS_VOIP)]
    df['Begin Time'] = pd.to_datetime(df['Begin Time'], errors='coerce')
    df.dropna(subset=['Begin Time'], inplace=True)
    df['day_key'] = df['Begin Time'].dt.date
    
    return df

# ==============================================================================
# FUNCIONES DE GENERACIÓN DE REPORTES
# ==============================================================================

def generar_reporte_agregado(df: pd.DataFrame, date_col_name: str) -> pd.DataFrame:
    """
    Genera estadísticas agregadas por extensión y día para el Reporte Isabel
    
    Args:
        df: DataFrame con datos de llamadas procesados
        date_col_name: Nombre de la columna que contiene las fechas
        
    Returns:
        pd.DataFrame: Reporte con totales, efectivas y última llamada por extensión
    """
    report_data = []
    
    for day_key in sorted(df['day_key'].unique()):
        day_data = df[df['day_key'] == day_key]
        
        extension_stats = day_data.groupby('Fuente').agg({
            date_col_name: ['count', 'max'],
            'Estado': lambda x: (x == ESTADO_EFECTIVA).sum()
        }).reset_index()
        
        extension_stats.columns = ['Fuente', 'Total_Llamadas', 'Ultima_Llamada', 'Llamadas_Efectivas']
        extension_stats['day_key'] = day_key
        
        report_data.append(extension_stats)
    
    if not report_data:
        return pd.DataFrame()
    
    final_report = pd.concat(report_data, ignore_index=True)
    final_report = final_report.rename(columns={
        'Fuente': 'Número Extensión',
        'Total_Llamadas': 'Total Llamadas',
        'Llamadas_Efectivas': 'Llamadas Efectivas',
        'Ultima_Llamada': 'Última Llamada'
    })
    
    # Formatear la columna de hora con segundos en formato de hora corta
    if 'Última Llamada' in final_report.columns:
        final_report['Última Llamada'] = pd.to_datetime(final_report['Última Llamada']).dt.strftime('%I:%M:%S %p')
    
    return final_report

def generar_reporte_agentes(df: pd.DataFrame) -> pd.DataFrame:
    """
    Genera estadísticas agregadas por agente y día para el Reporte VOIP
    Solo incluye: Extensión, Total (suma de todos los estados válidos), Última Llamada
    Estados válidos: answered, busy, no_answer, normal, out_area, offline
    
    Args:
        df: DataFrame con datos de agentes procesados y filtrados
        
    Returns:
        pd.DataFrame: Reporte con estadísticas de Total por agente
    """
    report_data = []
    
    for day_key in sorted(df['day_key'].unique()):
        day_data = df[df['day_key'] == day_key]
        
        # Agrupar por First Call Agent - solo contar total y última llamada
        agent_stats = day_data.groupby('First Call Agent').agg({
            'Ring Type': 'count',  # Total de llamadas válidas
            'Begin Time': 'max'    # Última llamada
        }).reset_index()
        
        # Renombrar columnas
        agent_stats.columns = ['Extensión', 'Total', 'Ultima Llamada']
        agent_stats['day_key'] = day_key
        
        report_data.append(agent_stats)
    
    if not report_data:
        return pd.DataFrame()
    
    final_report = pd.concat(report_data, ignore_index=True)
    
    # Formatear la columna de hora como HH:MM AM/PM
    if 'Ultima Llamada' in final_report.columns:
        final_report['Ultima Llamada'] = pd.to_datetime(final_report['Ultima Llamada']).dt.strftime('%I:%M %p')
    
    return final_report

# ==============================================================================
# FUNCIONES DE GENERACIÓN DE EXCEL
# ==============================================================================

def generar_excel_generico(report: pd.DataFrame, columnas_requeridas: List[str], 
                          columnas_config: Dict, tabla_prefix: str) -> io.BytesIO:
    """
    Genera archivo Excel genérico con formato de tabla dividido por días
    
    Args:
        report: DataFrame con el reporte
        columnas_requeridas: Lista de columnas que deben estar presentes
        columnas_config: Configuración de formato de columnas
        tabla_prefix: Prefijo para nombres de tabla
        
    Returns:
        BytesIO con el archivo Excel
    """
    output = io.BytesIO()
    
    if report.empty:
        raise ValueError("No hay datos para generar el reporte")
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        unique_days = sorted(report['day_key'].unique())
        
        if len(unique_days) == 0:
            raise ValueError("No hay fechas válidas en los datos")
        
        hojas_creadas = 0
        
        for day in unique_days:
            day_data = report[report['day_key'] == day]
            
            if day_data.empty:
                continue
            
            # Ordenar por primera columna si existe
            if len(columnas_requeridas) > 0 and columnas_requeridas[0] in day_data.columns:
                sheet_data = day_data.sort_values(by=columnas_requeridas[0])
            else:
                sheet_data = day_data
            
            # Verificar columnas requeridas
            if not all(col in sheet_data.columns for col in columnas_requeridas):
                continue
            
            final_report_df = sheet_data[columnas_requeridas]
            
            # Crear hoja
            sheet_name = generar_nombre_hoja_fecha(day)
            final_report_df.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            
            # Aplicar formato
            configurar_hoja_excel(worksheet, final_report_df, columnas_config, tabla_prefix)
            hojas_creadas += 1
        
        # Verificar que al menos una hoja fue creada
        if hojas_creadas == 0:
            # Crear una hoja de respaldo con información de error
            error_df = pd.DataFrame({
                'Error': ['No hay datos válidos para mostrar'],
                'Columnas Esperadas': [', '.join(columnas_requeridas)],
                'Datos Disponibles': [f"Filas: {len(report)}, Columnas: {', '.join(report.columns)}"]
            })
            error_df.to_excel(writer, sheet_name='Error_Info', index=False)
    
    output.seek(0)
    return output

# Funciones wrapper eliminadas - usar generar_excel_generico directamente

# ==============================================================================
# FUNCIONES DE NOMBRES DE ARCHIVO
# ==============================================================================

def generar_nombre_archivo_generico(report: pd.DataFrame, prefijo: str) -> str:
    """
    Genera nombre dinámico de archivo basado en las fechas de los datos
    
    Args:
        report: DataFrame con columna 'day_key' 
        prefijo: Prefijo del archivo (ej: "Reporte Llamadas Isabel")
        
    Returns:
        str: Nombre del archivo con formato legible
        
    Examples:
        - Un día: "Reporte Llamadas Isabel (9 Agosto 2025).xlsx"
        - Mismo mes: "Reporte Llamadas Isabel (3-8 Septiembre 2025).xlsx" 
        - Diferente mes: "Reporte Llamadas Isabel (30 Agosto - 5 Septiembre 2025).xlsx"
    """
    if report.empty or 'day_key' not in report.columns:
        timestamp = datetime.now().strftime("%Y-%m-%dT%H%M%S.%f")[:-3]
        return f"{prefijo} - {timestamp}.xlsx"
    
    fechas_unicas = sorted(report['day_key'].unique())
    fecha_inicio = fechas_unicas[0]
    fecha_fin = fechas_unicas[-1]
    
    # Un solo día
    if fecha_inicio == fecha_fin:
        dia = fecha_inicio.day
        mes = MESES_ESPANOL[fecha_inicio.month]
        año = fecha_inicio.year
        return f"{prefijo} ({dia} {mes} {año}).xlsx"
    
    # Múltiples días - analizar rangos
    dia_inicio, dia_fin = fecha_inicio.day, fecha_fin.day
    año_inicio, año_fin = fecha_inicio.year, fecha_fin.year
    mes_inicio_num, mes_fin_num = fecha_inicio.month, fecha_fin.month
    
    mes_inicio = MESES_ESPANOL[mes_inicio_num]
    mes_fin = MESES_ESPANOL[mes_fin_num]
    
    # Mismo mes y año
    if mes_inicio_num == mes_fin_num and año_inicio == año_fin:
        return f"{prefijo} ({dia_inicio}-{dia_fin} {mes_inicio} {año_inicio}).xlsx"
    
    # Mismo año, diferente mes  
    elif año_inicio == año_fin:
        return f"{prefijo} ({dia_inicio} {mes_inicio} - {dia_fin} {mes_fin} {año_inicio}).xlsx"
    
    # Diferente año
    else:
        return f"{prefijo} ({dia_inicio} {mes_inicio} {año_inicio} - {dia_fin} {mes_fin} {año_fin}).xlsx"

# ==============================================================================
# FUNCIONES PRINCIPALES
# ==============================================================================

def procesar_llamadas_isabel():
    """
    Procesa archivos CSV o Excel de llamadas y genera el Reporte Isabel
    """
    try:
        # Validar archivos de entrada
        archivos, error_msg, status_code = validar_archivos_entrada()
        if error_msg:
            return error_msg, status_code
        
        # Procesar datos (CSV o Excel, acepta múltiples archivos)
        df, date_col_name = procesar_datos_csv(*archivos)
        
        # Generar reporte agregado
        report = generar_reporte_agregado(df, date_col_name)
        
        # Generar nombre dinámico del archivo
        nombre_archivo = generar_nombre_archivo_generico(report, "Reporte Llamadas Isabel")
        
        # Generar archivo Excel
        columnas_requeridas = ['Número Extensión', 'Total Llamadas', 'Llamadas Efectivas', 'Última Llamada']
        output = generar_excel_generico(report, columnas_requeridas, COLUMNAS_CONFIG_ISABEL, "Tabla")
        
        # Resetear el buffer para envío
        output.seek(0)
        
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=nombre_archivo)

    except ValueError as ve:
        return f"Error de validación: {ve}", 400
    except Exception as e:
        return f"Ocurrió un error procesando los archivos: {e}", 500

def procesar_reporte_agentes():
    """
    Función principal para procesar reporte VOIP
    """
    try:
        # Validar archivos
        archivos, error_msg, status_code = validar_archivos_entrada()
        if error_msg:
            return error_msg, status_code
        
        # Procesar datos
        df = procesar_datos_agentes(*archivos)
        
        if df.empty:
            return "Error: No hay datos válidos después del filtrado. Verifique que los archivos contengan llamadas con estados válidos (answered, busy, no_answer, normal, out_area, offline).", 400
        
        # Generar reporte
        report = generar_reporte_agentes(df)
        
        if report.empty:
            return "Error: No se pudo generar el reporte. Verifique que los datos contengan agentes válidos.", 400
        
        # Generar nombre dinámico del archivo
        nombre_archivo = generar_nombre_archivo_generico(report, "Reporte Llamadas VOIP")
        
        # Generar Excel
        columnas_requeridas = ['Extensión', 'Total', 'Ultima Llamada']
        output = generar_excel_generico(report, columnas_requeridas, COLUMNAS_CONFIG_VOIP, "Agentes")
        
        # Resetear el buffer para envío
        output.seek(0)
        
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        as_attachment=True, download_name=nombre_archivo)

    except ValueError as ve:
        return f"Error de validación: {ve}", 400
    except Exception as e:
        return f"Ocurrió un error procesando los archivos: {e}", 500